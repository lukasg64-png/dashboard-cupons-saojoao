import os
import json
import re
import unicodedata
import pandas as pd
import numpy as np
from datetime import datetime, timezone, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = r'c:\Users\lucas.alves6\OneDrive - Farmácias São João\Documentos\ANTIGRAVITI'
HIERARCHY_EXCEL_PATH = os.path.join(BASE_DIR, 'Dashboard desempenho Diretoria', 'base Lojas Atualizado.xlsx')
CACHE_PATH = os.path.join(BASE_DIR, 'Dashboard desempenho Diretoria', 'dashboard-cupons-unificado', 'server', 'data', 'vtex_orders_cache.json')
OUTPUT_EXCEL_PATH = os.path.join(BASE_DIR, 'Relatorio_Uso_Cupons_Agosto_2026_Completo.xlsx')
OUTPUT_EXCEL_COPY_PATH = os.path.join(BASE_DIR, 'Dashboard desempenho Diretoria', 'Relatorio_Uso_Cupons_Agosto_2026_Completo.xlsx')

# Fuso Horário Oficial de Brasília (UTC-3)
BRT_TIMEZONE = timezone(timedelta(hours=-3))

def parse_to_brasilia_time(iso_str):
    if not iso_str:
        return None, None
    try:
        dt_utc = datetime.fromisoformat(str(iso_str).replace('Z', '+00:00'))
        dt_brt = dt_utc.astimezone(BRT_TIMEZONE)
        return dt_brt.strftime('%Y-%m-%d'), dt_brt.strftime('%H:%M:%S')
    except Exception:
        return None, None

ABBREVIATION_MAP = {
    'baln': 'balneario', 'bal': 'balneario', 'floripa': 'florianopolis',
    'sta': 'santa', 'sto': 'santo', 'eng': 'engenheiro', 'mal': 'marechal',
    'dioni': 'dionisio', 'cnel': 'coronel', 'cel': 'coronel',
    'fco': 'francisco', 'franc': 'francisco', 'gal': 'galeria',
    'hosp': 'hospital', 'louren': 'lourenco', 'terez': 'terezinha',
    'ant': 'antonio', 's': 'sao', 'dr': 'doutor', 'av': 'avenida',
    'gen': 'general', 'v': 'vila', 'vl': 'vila', 'sn': 'santo',
    'st': 'santo', 'pt': 'ponto', 'pto': 'porto', 'distr': 'distrito',
    'pres': 'presidente',
}

CITY_SUFFIX_MAP = {
    'sapucaia': 'sapucaia sul', 'venancio': 'venancio aires',
    'rosario': 'rosario do sul', 'cachoeira': 'cachoeira do sul',
    'sao lourenco do sul': 'sao lourenco', 'sao lourenco oeste': 'sao lourenco do oeste',
    'sao lourenco do oeste': 'sao lourenco oeste', 'sao sebastiao cai': 'sao sebastiao',
    'julio castilhos': 'julio de castilhos', 'quedas iguacu': 'quedas do iguacu',
    'cruzeiro oeste': 'cruzeiro do oeste', 'sao miguel iguacu': 'sao miguel do iguacu',
    'encruzilhada sul': 'encruzilhada do sul', 'cerro grande sul': 'cerro grande',
    'cerro grande do sul': 'cerro grande', 'sao miguel oeste': 'sao miguel do oeste',
    'bela vista paraiso': 'bela vista do paraiso',
    'balneario arroio silva': 'balneario arroio do silva',
    'sao pedro sul': 'sao pedro do sul', 'sao francisco de assis': 'sao francisco assis',
    'sao francisco assis': 'sao francisco assis',
    'santana livramento': 'santana do livramento',
    'santa vitoria palmar': 'santa vitoria do palmar',
    'sao jose norte': 'sao jose do norte', 'sao joao do oeste': 'sao joao oeste',
    'sao luiz gonzaga': 'sao luiz', 'sao marcos': 'sao marcos',
    'herval d oeste': 'herval doeste', 'herval doeste': 'herval doeste',
}

SPECIAL_VTEX_TO_CSV = {
    'farmacias sao joao delivery': 'porto alegre dark store',
    'sjdigital1601': 'santo antonio das missoes',
    'pf': 'pf matriz', 'pf matriz': 'pf matriz', 'pf modelo': 'pf loja modelo',
    'pf uruguai': 'pf uruguai', 'pf shopping bella': 'pf shopping',
    'pf general netto': 'pf general neto', 'gruarapuava': 'guarapuava',
    'santo amaro': 'santo amaro imperatriz', 'sao francisco paula': 'sao fran paula',
    'sao francisco de paula': 'sao fran paula',
    'santa terezinha de itaipu': 'santa terezinha do itaipu',
    'santa terezinha itaipu': 'santa terezinha do itaipu',
    'sta terez de itaipu': 'santa terezinha do itaipu',
    'santo antonio missoes': 'santo antonio das missoes',
    'caxias 21': 'caxias 20',
}

def normalize_text(text):
    if not text: return ''
    s = str(text).lower()
    s = unicodedata.normalize('NFD', s)
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    s = re.sub(r'[^a-z0-9]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s

def clean_vtex_seller(seller_name):
    if not seller_name: return ''
    cleaned = re.sub(r'\s*-\s*[\d\.\/\-]{11,25}\s*-\s*\d+\s*$', '', seller_name, flags=re.IGNORECASE).strip()
    if cleaned == seller_name and ' - ' in seller_name:
        parts = seller_name.split(' - ')
        if len(parts) >= 3 and re.match(r'^\d+$', parts[-1].strip()):
            cleaned = ' - '.join(parts[:-2]).strip()
        else:
            cleaned = parts[0].strip()
    return cleaned

def canonicalize(norm_name):
    res = str(norm_name).lower()
    if res in SPECIAL_VTEX_TO_CSV and SPECIAL_VTEX_TO_CSV[res] != res:
        return canonicalize(SPECIAL_VTEX_TO_CSV[res])
    res = re.sub(r'([a-z])(\d)', r'\1 \2', res)
    res = re.sub(r'\b0+(\d+)\b', r'\1', res)
    res = re.sub(r'\s+(rs|pr|sc)\s*$', '', res)
    res = re.sub(r'\s+(rs|pr|sc)\s+(\d)', r' \2', res)
    res = re.sub(r'\s*-\s*(nova|shop|gal|hosp|merc|pr|sc|rs)\b', '', res, flags=re.IGNORECASE)
    res = re.sub(r'\b(nova|shop|gal|hosp|merc)\b', '', res, flags=re.IGNORECASE)
    res = re.sub(r'\bnv\b|\bnov\b|\b1nov\b', '', res)
    res = re.sub(r'\s+', ' ', res).strip()

    words = res.split(' ')
    expanded = [ABBREVIATION_MAP.get(w, w) for w in words]
    res = ' '.join(expanded)
    res = re.sub(r'\bd\s+', 'd', res)
    res = re.sub(r"\bd'", 'd', res)

    m = re.match(r'^(.+?)\s+(\d+)$', res)
    if m:
        base_name, num = m.group(1).strip(), m.group(2)
        if base_name in CITY_SUFFIX_MAP:
            res = f"{CITY_SUFFIX_MAP[base_name]} {num}"
    else:
        if res in CITY_SUFFIX_MAP:
            res = CITY_SUFFIX_MAP[res]

    m_final = re.match(r'^(.+?)\s+(\d+)$', res)
    if m_final:
        b_name = m_final.group(1).strip()
        if b_name in SPECIAL_VTEX_TO_CSV and SPECIAL_VTEX_TO_CSV[b_name] != b_name:
            res = f"{SPECIAL_VTEX_TO_CSV[b_name]} {m_final.group(2)}"
    if res in SPECIAL_VTEX_TO_CSV and SPECIAL_VTEX_TO_CSV[res] != res:
        return canonicalize(SPECIAL_VTEX_TO_CSV[res])
    return re.sub(r'\s+', ' ', res).strip()

class StoreMatcher:
    def __init__(self, excel_path):
        self.df_stores = pd.read_excel(excel_path)
        self.stores_map = {}
        self.canon_map = {}
        self.base_no_num_map = {}
        self.registered_stores = []

        for _, row in self.df_stores.iterrows():
            store_name = str(row['Desc_Filial']).strip()
            norm = normalize_text(store_name)
            canon = canonicalize(norm)
            
            info = {
                'Desc_Filial': store_name,
                'Diretor': str(row['Diretor']).strip(),
                'Distrital': str(row['Distrital']).strip(),
                'Coordenador': str(row['Coordenador']).strip()
            }
            self.registered_stores.append(info)
            self.stores_map[store_name] = info
            self.stores_map[norm] = info
            self.canon_map[canon] = info
            self.canon_map[norm] = info

            if not re.search(r'\b\d+\b', canon):
                self.base_no_num_map[canon] = info
                self.canon_map[canon + ' 1'] = info
            else:
                m1 = re.match(r'^(.+?)\s+1$', canon)
                if m1:
                    self.base_no_num_map[m1.group(1).strip()] = info

    def match(self, raw_seller_name):
        if not raw_seller_name:
            return None
        cleaned = clean_vtex_seller(raw_seller_name)
        norm = normalize_text(cleaned)

        if cleaned in self.stores_map:
            return self.stores_map[cleaned]
        if norm in self.stores_map:
            return self.stores_map[norm]
        if norm in self.canon_map:
            return self.canon_map[norm]

        canon = canonicalize(norm)
        if canon in self.canon_map:
            return self.canon_map[canon]

        if canon.endswith(' 1'):
            base = canon[:-2].strip()
            if base in self.base_no_num_map:
                return self.base_no_num_map[base]
            if base in self.canon_map:
                return self.canon_map[base]

        if not re.search(r'\b\d+\b', canon):
            with1 = canon + ' 1'
            if with1 in self.canon_map:
                return self.canon_map[with1]

        for s_canon, info in self.canon_map.items():
            if len(s_canon) > 4 and (s_canon in canon or canon in s_canon):
                return info

        return None

def build_excel_report():
    print("[1/5] Carregando cadastro de filiais e hierarquia corporativa...")
    matcher = StoreMatcher(HIERARCHY_EXCEL_PATH)
    print(f"Total de lojas no cadastro: {len(matcher.registered_stores)}")

    print("[2/5] Carregando e convertendo pedidos da VTEX para o Horário de Brasília (UTC-3)...")
    with open(CACHE_PATH, 'r', encoding='utf-8') as f:
        cache = json.load(f)
    print(f"Total de pedidos no cache: {len(cache)}")

    # Filtra mês completo de Agosto/2026 (01/08/2026 00:00:00 BRT até 31/08/2026 23:59:59 BRT)
    august_orders_with_coupon = []
    coupon_counts_global = {}
    total_august_orders = 0
    dates_found = set()

    for oid, o in cache.items():
        if not isinstance(o, dict): continue
        cd = o.get('creationDate')
        if not cd: continue
        
        # Converte para Horário de Brasília
        date_brt, time_brt = parse_to_brasilia_time(cd)
        if not date_brt: continue

        if '2026-08-01' <= date_brt <= '2026-08-31':
            total_august_orders += 1
            dates_found.add(date_brt)
            
            raw_coupon = o.get('coupon')
            if not raw_coupon:
                m = o.get('marketingData') or {}
                raw_coupon = m.get('coupon')
            if not raw_coupon:
                rates = o.get('ratesAndBenefitsData') or {}
                for c in rates.get('coupon', []):
                    if c:
                        raw_coupon = c
                        break
            
            if raw_coupon:
                coupon_clean = str(raw_coupon).strip().upper()
                if coupon_clean:
                    seller_name = o.get('sellers', [{}])[0].get('name', '') if o.get('sellers') else ''
                    matched = matcher.match(seller_name)
                    
                    val_brl = (o.get('value') or 0) / 100.0
                    
                    august_orders_with_coupon.append({
                        'orderId': oid,
                        'date_brt': date_brt,
                        'time_brt': time_brt,
                        'coupon': coupon_clean,
                        'seller_raw': seller_name,
                        'loja': matched['Desc_Filial'] if matched else (clean_vtex_seller(seller_name) or 'E-COMMERCE CENTRAL / NÃO MAPEADA'),
                        'coordenador': matched['Coordenador'] if matched else 'NÃO MAPEADO',
                        'distrital': matched['Distrital'] if matched else 'NÃO MAPEADO',
                        'diretor': matched['Diretor'] if matched else 'NÃO MAPEADO',
                        'valor': val_brl
                    })
                    coupon_counts_global[coupon_clean] = coupon_counts_global.get(coupon_clean, 0) + 1

    print(f"Dias no Horário de Brasília: {len(dates_found)} dias ({min(dates_found)} a {max(dates_found)})")
    print(f"Total de pedidos em Agosto (BRT): {total_august_orders:,}")
    print(f"Total de pedidos com cupom em Agosto (BRT): {len(august_orders_with_coupon):,}")
    print(f"Total de cupons distintos utilizados: {len(coupon_counts_global)}")

    # Ordena os cupons pelo volume total de uso decrescente
    sorted_coupons = sorted(coupon_counts_global.keys(), key=lambda c: coupon_counts_global[c], reverse=True)

    print("[3/5] Estruturando matriz por Filial, Coordenação, Distrital e Diretoria...")
    
    # Inicializa todas as lojas do cadastro da rede
    stores_data = {}
    for st in matcher.registered_stores:
        key = (st['Desc_Filial'], st['Coordenador'], st['Distrital'], st['Diretor'])
        stores_data[key] = {c: 0 for c in sorted_coupons}

    # Preenche a contagem de uso de cupons por filial
    for ord_item in august_orders_with_coupon:
        key = (ord_item['loja'], ord_item['coordenador'], ord_item['distrital'], ord_item['diretor'])
        if key not in stores_data:
            stores_data[key] = {c: 0 for c in sorted_coupons}
        stores_data[key][ord_item['coupon']] += 1

    # Converte em linhas para o DataFrame
    rows = []
    for (filial, coord, dist, dir_nome), cupom_dict in stores_data.items():
        total_loja = sum(cupom_dict.values())
        row = {
            'Filial': filial,
            'Coordenação': coord,
            'Distrital': dist,
            'Diretoria': dir_nome,
            'Total de Cupons': total_loja
        }
        for c in sorted_coupons:
            row[c] = cupom_dict[c]
        rows.append(row)

    df_matrix = pd.DataFrame(rows)
    # Ordena por Diretoria, Distrital, Coordenação, Total de Cupons decrescente
    df_matrix = df_matrix.sort_values(by=['Diretoria', 'Distrital', 'Coordenação', 'Total de Cupons', 'Filial'], ascending=[True, True, True, False, True])

    print("[4/5] Gerando resumos analíticos adicionais...")
    # Resumo por Cupom
    coupon_summary_rows = []
    df_coupons_raw = pd.DataFrame(august_orders_with_coupon)
    for c in sorted_coupons:
        sub = df_coupons_raw[df_coupons_raw['coupon'] == c]
        coupon_summary_rows.append({
            'Cupom': c,
            'Quantidade de Usos': len(sub),
            '% do Total': len(sub) / len(august_orders_with_coupon),
            'Faturamento Gerado (R$)': sub['valor'].sum(),
            'Ticket Médio (R$)': sub['valor'].mean() if len(sub) > 0 else 0,
            'Filiais Únicas Atendidas': sub['loja'].nunique()
        })
    df_coupon_summary = pd.DataFrame(coupon_summary_rows)

    # Resumo por Diretoria & Distrital
    df_hier_summary = df_matrix.groupby(['Diretoria', 'Distrital']).agg(
        Total_Lojas=('Filial', 'count'),
        Lojas_Com_Uso=('Total de Cupons', lambda x: (x > 0).sum()),
        Total_Cupons_Usados=('Total de Cupons', 'sum')
    ).reset_index()

    print("[5/5] Formatando Excel profissionalmente com openpyxl...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Cores e Estilos Corporativos
    header_fill = PatternFill(start_color="0B3060", end_color="0B3060", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    total_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    align_left = Alignment(horizontal='left', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # ==========================================
    # ABA 1: Uso de Cupons por Filial
    # ==========================================
    ws1 = wb.create_sheet(title="Uso de Cupons por Filial")
    ws1.views.sheetView[0].showGridLines = True

    # Escreve cabeçalhos
    headers_ws1 = list(df_matrix.columns)
    for col_num, h_name in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=1, column=col_num, value=h_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center if col_num > 4 else align_left
        cell.border = thin_border
    ws1.row_dimensions[1].height = 28

    # Escreve dados
    for row_idx, row_data in enumerate(df_matrix.itertuples(index=False), 2):
        ws1.row_dimensions[row_idx].height = 19
        is_even = (row_idx % 2 == 0)
        row_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid") if is_even else None

        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = bold_font if col_idx == 5 else regular_font
            cell.border = thin_border
            if row_fill and col_idx != 5:
                cell.fill = row_fill
            elif col_idx == 5:
                cell.fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
            
            if col_idx in [1, 2, 3, 4]:
                cell.alignment = align_left
            else:
                cell.alignment = align_right
                cell.number_format = '#,##0'

    # Linha de Total Geral no rodapé
    last_row = len(df_matrix) + 2
    ws1.row_dimensions[last_row].height = 24
    cell_tot_label = ws1.cell(row=last_row, column=1, value="TOTAL GERAL DA REDE")
    cell_tot_label.font = total_font
    cell_tot_label.fill = total_fill
    cell_tot_label.alignment = align_left
    cell_tot_label.border = thin_border

    for c_idx in range(2, 5):
        c_empty = ws1.cell(row=last_row, column=c_idx, value="")
        c_empty.font = total_font
        c_empty.fill = total_fill
        c_empty.border = thin_border

    # Fórmulas de soma para cada coluna numérica
    for c_idx in range(5, len(headers_ws1) + 1):
        col_letter = get_column_letter(c_idx)
        c_sum = ws1.cell(row=last_row, column=c_idx, value=f"=SUM({col_letter}2:{col_letter}{last_row-1})")
        c_sum.font = total_font
        c_sum.fill = total_fill
        c_sum.alignment = align_right
        c_sum.border = thin_border
        c_sum.number_format = '#,##0'

    # Congela painéis: Linha 1 e Colunas A a E
    ws1.freeze_panes = 'F2'
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers_ws1))}{last_row-1}"

    # Auto-ajuste de largura de colunas
    for col_idx in range(1, len(headers_ws1) + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws1.column_dimensions[col_letter].width = 30
        elif col_idx in [2, 3]:
            ws1.column_dimensions[col_letter].width = 24
        elif col_idx == 4:
            ws1.column_dimensions[col_letter].width = 18
        elif col_idx == 5:
            ws1.column_dimensions[col_letter].width = 16
        else:
            col_name = headers_ws1[col_idx - 1]
            ws1.column_dimensions[col_letter].width = max(len(col_name) + 3, 12)

    # ==========================================
    # ABA 2: Resumo por Cupom
    # ==========================================
    ws2 = wb.create_sheet(title="Ranking de Cupons")
    ws2.views.sheetView[0].showGridLines = True
    ws2.row_dimensions[1].height = 28

    headers_ws2 = list(df_coupon_summary.columns)
    for col_num, h_name in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=1, column=col_num, value=h_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center if col_num > 1 else align_left
        cell.border = thin_border

    for r_idx, r_data in enumerate(df_coupon_summary.itertuples(index=False), 2):
        ws2.row_dimensions[r_idx].height = 19
        is_even = (r_idx % 2 == 0)
        row_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid") if is_even else None
        
        for c_idx, val in enumerate(r_data, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if row_fill: cell.fill = row_fill
            
            if c_idx == 1:
                cell.alignment = align_left
                cell.font = bold_font
            elif c_idx in [2, 6]:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif c_idx == 3:
                cell.alignment = align_right
                cell.number_format = '0.00%'
            elif c_idx in [4, 5]:
                cell.alignment = align_right
                cell.number_format = 'R$ #,##0.00'

    # Total rodapé Aba 2
    last_r2 = len(df_coupon_summary) + 2
    ws2.row_dimensions[last_r2].height = 24
    c2_tot = ws2.cell(row=last_r2, column=1, value="TOTAL GERAL")
    c2_tot.font = total_font
    c2_tot.fill = total_fill
    c2_tot.border = thin_border

    c2_sum_q = ws2.cell(row=last_r2, column=2, value=f"=SUM(B2:B{last_r2-1})")
    c2_sum_q.font = total_font
    c2_sum_q.fill = total_fill
    c2_sum_q.border = thin_border
    c2_sum_q.number_format = '#,##0'

    c2_sum_pct = ws2.cell(row=last_r2, column=3, value=f"=SUM(C2:C{last_r2-1})")
    c2_sum_pct.font = total_font
    c2_sum_pct.fill = total_fill
    c2_sum_pct.border = thin_border
    c2_sum_pct.number_format = '0.00%'

    c2_sum_fat = ws2.cell(row=last_r2, column=4, value=f"=SUM(D2:D{last_r2-1})")
    c2_sum_fat.font = total_font
    c2_sum_fat.fill = total_fill
    c2_sum_fat.border = thin_border
    c2_sum_fat.number_format = 'R$ #,##0.00'

    c2_tk_avg = ws2.cell(row=last_r2, column=5, value=f"=D{last_r2}/B{last_r2}")
    c2_tk_avg.font = total_font
    c2_tk_avg.fill = total_fill
    c2_tk_avg.border = thin_border
    c2_tk_avg.number_format = 'R$ #,##0.00'

    c2_lojas = ws2.cell(row=last_r2, column=6, value="")
    c2_lojas.font = total_font
    c2_lojas.fill = total_fill
    c2_lojas.border = thin_border

    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = f"A1:F{last_r2-1}"
    for col_idx in range(1, 7):
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = [25, 20, 15, 24, 18, 24][col_idx - 1]

    # ==========================================
    # ABA 3: Resumo por Diretoria e Distrital
    # ==========================================
    ws3 = wb.create_sheet(title="Resumo por Diretoria")
    ws3.views.sheetView[0].showGridLines = True
    ws3.row_dimensions[1].height = 28

    headers_ws3 = ['Diretoria', 'Distrital', 'Total de Lojas', 'Lojas com Uso de Cupons', '% Cobertura', 'Total de Cupons Usados', 'Média Cupons/Loja']
    for col_num, h_name in enumerate(headers_ws3, 1):
        cell = ws3.cell(row=1, column=col_num, value=h_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center if col_num > 2 else align_left
        cell.border = thin_border

    for r_idx, r_data in enumerate(df_hier_summary.itertuples(index=False), 2):
        ws3.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)
        row_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid") if is_even else None
        
        diretoria, distrital, tot_lojas, lojas_uso, tot_cupons = r_data
        cob_pct = lojas_uso / tot_lojas if tot_lojas > 0 else 0
        med_cup = tot_cupons / tot_lojas if tot_lojas > 0 else 0

        row_vals = [diretoria, distrital, tot_lojas, lojas_uso, cob_pct, tot_cupons, med_cup]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if row_fill: cell.fill = row_fill
            
            if c_idx in [1, 2]:
                cell.alignment = align_left
                if c_idx == 1: cell.font = bold_font
            elif c_idx in [3, 4, 6]:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif c_idx == 5:
                cell.alignment = align_right
                cell.number_format = '0.0%'
            elif c_idx == 7:
                cell.alignment = align_right
                cell.number_format = '#,##0.0'

    # Total rodapé Aba 3
    last_r3 = len(df_hier_summary) + 2
    ws3.row_dimensions[last_r3].height = 24
    c3_tot = ws3.cell(row=last_r3, column=1, value="TOTAL GERAL")
    c3_tot.font = total_font
    c3_tot.fill = total_fill
    c3_tot.border = thin_border

    c3_blank = ws3.cell(row=last_r3, column=2, value="")
    c3_blank.font = total_font
    c3_blank.fill = total_fill
    c3_blank.border = thin_border

    c3_tot_l = ws3.cell(row=last_r3, column=3, value=f"=SUM(C2:C{last_r3-1})")
    c3_tot_l.font = total_font
    c3_tot_l.fill = total_fill
    c3_tot_l.border = thin_border
    c3_tot_l.number_format = '#,##0'

    c3_tot_u = ws3.cell(row=last_r3, column=4, value=f"=SUM(D2:D{last_r3-1})")
    c3_tot_u.font = total_font
    c3_tot_u.fill = total_fill
    c3_tot_u.border = thin_border
    c3_tot_u.number_format = '#,##0'

    c3_tot_cob = ws3.cell(row=last_r3, column=5, value=f"=D{last_r3}/C{last_r3}")
    c3_tot_cob.font = total_font
    c3_tot_cob.fill = total_fill
    c3_tot_cob.border = thin_border
    c3_tot_cob.number_format = '0.0%'

    c3_tot_c = ws3.cell(row=last_r3, column=6, value=f"=SUM(F2:F{last_r3-1})")
    c3_tot_c.font = total_font
    c3_tot_c.fill = total_fill
    c3_tot_c.border = thin_border
    c3_tot_c.number_format = '#,##0'

    c3_tot_med = ws3.cell(row=last_r3, column=7, value=f"=F{last_r3}/C{last_r3}")
    c3_tot_med.font = total_font
    c3_tot_med.fill = total_fill
    c3_tot_med.border = thin_border
    c3_tot_med.number_format = '#,##0.0'

    ws3.freeze_panes = 'A2'
    ws3.auto_filter.ref = f"A1:G{last_r3-1}"
    for col_idx in range(1, 8):
        col_letter = get_column_letter(col_idx)
        ws3.column_dimensions[col_letter].width = [20, 24, 16, 24, 16, 24, 20][col_idx - 1]

    wb.save(OUTPUT_EXCEL_PATH)
    wb.save(OUTPUT_EXCEL_COPY_PATH)
    print(f"\n[SUCESSO] Relatório Excel gerado com sucesso!")
    print(f"Salvo em: {OUTPUT_EXCEL_PATH}")
    print(f"Salvo cópia em: {OUTPUT_EXCEL_COPY_PATH}")

    return {
        'total_pedidos_agosto': total_august_orders,
        'total_cupons_agosto': len(august_orders_with_coupon),
        'total_cupons_distintos': len(sorted_coupons),
        'total_lojas_matriz': len(df_matrix),
        'top_10_cupons': df_coupon_summary.head(10).to_dict(orient='records'),
        'por_diretoria': df_matrix.groupby('Diretoria')['Total de Cupons'].sum().to_dict(),
        'dias_cobertos': f"{min(dates_found)} a {max(dates_found)} ({len(dates_found)} dias)"
    }

if __name__ == '__main__':
    stats = build_excel_report()
    print("\n--- RESUMO DE AGOSTO (HORÁRIO DE BRASÍLIA UTC-3) ---")
    print(f"Período: {stats['dias_cobertos']}")
    print(f"Total de Pedidos: {stats['total_pedidos_agosto']:,}")
    print(f"Total de Usos de Cupons: {stats['total_cupons_agosto']:,}")
    print(f"Cupons Únicos: {stats['total_cupons_distintos']}")
    print(f"Total de Lojas Mapeadas: {stats['total_lojas_matriz']}")
    print("\nCupons por Diretoria:")
    for d, val in stats['por_diretoria'].items():
        print(f"  - {d}: {val:,} cupons")
